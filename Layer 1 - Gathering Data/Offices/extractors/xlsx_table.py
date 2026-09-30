"""Structured spreadsheet extractor (tier B).

Layouts are declared per doc_key under documents.yaml → table.
Until a layout is filled in (and preferably approved once via review),
fetch raises AdapterError so the planner marks the source as error/overdue
rather than inventing numbers.
"""

from __future__ import annotations

import logging
import re
from datetime import date
from pathlib import Path
from typing import Any

from adapters.base import AdapterError
from doc_cache import CachedDocument
from documents_registry import DocSpec
from models import Observation
from registry import Entry

log = logging.getLogger(__name__)


def _open_workbook(path: Path):
    try:
        import openpyxl
    except ImportError as exc:
        raise AdapterError(
            "openpyxl is required for XLSX extraction — pip install openpyxl"
        ) from exc
    return openpyxl.load_workbook(path, data_only=True, read_only=True)


def _parse_period_cell(raw: Any, frequency: str) -> date | None:
    if raw is None:
        return None
    if hasattr(raw, "year") and hasattr(raw, "month"):
        d = raw
        return date(d.year, 1, 1) if frequency == "yearly" else date(d.year, d.month, 1)
    s = str(raw).strip()
    m = re.fullmatch(r"(20\d{2})(?:[-/.](\d{1,2}))?", s)
    if not m:
        return None
    year = int(m.group(1))
    month = int(m.group(2) or 1)
    if frequency == "yearly":
        return date(year, 1, 1)
    if not 1 <= month <= 12:
        return None
    return date(year, month, 1)


def extract_xlsx_table(
    spec: DocSpec,
    doc: CachedDocument,
    entries: list[Entry],
    *,
    since: date,
    today: date,
) -> list[Observation]:
    table = spec.table or {}
    date_column = table.get("date_column")
    value_columns = table.get("value_columns") or {}
    if date_column is None or not value_columns:
        raise AdapterError(
            f"xlsx_table[{spec.key}]: table.date_column / value_columns not configured yet "
            f"(capture a fixture workbook and fill documents.yaml)"
        )

    # Map registry key → column name declared in documents.yaml
    key_to_col: dict[int, str] = {}
    for e in entries:
        meta = spec.indexes.get(e.key) or {}
        field = str(meta.get("field") or e.params.get("field") or e.params.get("row") or "")
        if field and field in value_columns:
            key_to_col[e.key] = str(value_columns[field])
        elif str(e.key) in value_columns:
            key_to_col[e.key] = str(value_columns[str(e.key)])
        elif e.key in value_columns:
            key_to_col[e.key] = str(value_columns[e.key])
    if not key_to_col:
        raise AdapterError(f"xlsx_table[{spec.key}]: no value columns matched due indexes")

    wb = _open_workbook(doc.path)
    try:
        sheet_id = table.get("sheet", 0)
        ws = wb[wb.sheetnames[sheet_id]] if isinstance(sheet_id, int) else wb[str(sheet_id)]
        header_row = int(table.get("header_row") or 0)
        rows = list(ws.iter_rows(values_only=True))
        if header_row >= len(rows):
            raise AdapterError(f"xlsx_table[{spec.key}]: header_row {header_row} out of range")

        # Resolve column letters/names → indexes
        headers = [str(c).strip() if c is not None else "" for c in rows[header_row]]

        def col_index(ref: Any) -> int:
            if isinstance(ref, int):
                return ref
            name = str(ref)
            if name in headers:
                return headers.index(name)
            # Excel letter
            if re.fullmatch(r"[A-Za-z]+", name):
                n = 0
                for ch in name.upper():
                    n = n * 26 + (ord(ch) - 64)
                return n - 1
            raise AdapterError(f"xlsx_table[{spec.key}]: unknown column {ref!r}")

        date_i = col_index(date_column)
        col_is = {k: col_index(c) for k, c in key_to_col.items()}
        out: list[Observation] = []
        for row in rows[header_row + 1 :]:
            if not row or date_i >= len(row):
                continue
            # Use the first entry's frequency; doc specs are single-frequency.
            freq = entries[0].frequency if entries else spec.frequency
            period = _parse_period_cell(row[date_i], freq)
            if period is None or period < since or period > today:
                continue
            for key, ci in col_is.items():
                if ci >= len(row) or row[ci] is None or row[ci] == "":
                    continue
                try:
                    value = float(str(row[ci]).replace(",", "").replace("%", "").strip())
                except ValueError:
                    continue
                raw = str(row[ci])
                out.append(
                    Observation(
                        key=key,
                        period=period,
                        value=value,
                        raw_value=raw,
                        source_url=doc.url,
                        method="table",
                        confidence=1.0,
                        evidence={
                            "doc_key": spec.key,
                            "document_sha256": doc.sha256,
                            "document_id": doc.db_id,
                            "edition": doc.edition,
                            "extractor": "xlsx_table",
                            "column": key_to_col[key],
                        },
                    )
                )
        return out
    finally:
        wb.close()
